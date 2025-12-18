# from pptx import Presentation
# from pptx.util import Inches, Pt
# from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # ✅ ADDED MSO_ANCHOR
# from pptx.dml.color import RGBColor
# from pptx.enum.shapes import MSO_SHAPE
# import json
# import os

# class JSONAdapter:
#     @staticmethod
#     def adapt_json_for_ppt(input_json_path, output_json_path):
#         """Adapt the image retrieval JSON structure to match PPT generator expectations"""
#         with open(input_json_path, 'r', encoding='utf-8') as f:
#             data = json.load(f)
        
#         # Transform the structure
#         adapted_data = {
#             "presentation_title": data.get("presentation_title", "AI Presentation"),
#             "slides": []
#         }
        
#         for slide in data.get("slides", []):
#             # Convert content to list format expected by PPT generator
#             content_list = []
#             if slide.get("content"):
#                 content_text = slide["content"]
#                 if isinstance(content_text, str):
#                     # Split content by sentences or newlines for bullet points
#                     sentences = [s.strip() for s in content_text.split('\n') if s.strip()]
#                     content_list = sentences[:6]  # Limit to 6 bullet points
#                 else:
#                     content_list = content_text
            
#             # Adapt images structure
#             suggested_images = []
#             for img in slide.get("images", []):
#                 suggested_images.append({
#                     "image_path": img.get("image_path", ""),
#                     "caption": img.get("caption", "Image"),
#                     "similarity_score": img.get("similarity_score", 0)
#                 })
            
#             adapted_slide = {
#                 "slide_number": slide.get("slide_number", len(adapted_data["slides"]) + 1),
#                 "slide_title": slide.get("slide_title", "Untitled Slide"),
#                 "content": content_list,
#                 "suggested_images": suggested_images
#             }
            
#             adapted_data["slides"].append(adapted_slide)
        
#         # Save adapted JSON
#         with open(output_json_path, 'w', encoding='utf-8') as f:
#             json.dump(adapted_data, f, indent=2)
        
#         print(f"✅ JSON adapted successfully for template2!")
#         return adapted_data

# class DynamicLayoutPresentationGenerator:
#     def __init__(self, json_file_path):
#         self.json_file_path = json_file_path
#         self.presentation_data = self.load_json_data()
#         self.prs = None
        
#         # 🌤️ LIGHT THEME COLORS
#         self.colors = {
#             'background': RGBColor(255, 255, 255),
#             'card_bg': RGBColor(248, 249, 250),
#             'primary': RGBColor(0, 102, 204),
#             'accent': RGBColor(255, 87, 34),
#             'success': RGBColor(46, 204, 113),
#             'text_primary': RGBColor(33, 37, 41),
#             'text_secondary': RGBColor(108, 117, 125),
#             'border': RGBColor(222, 226, 230),
#             'image_border': RGBColor(0, 102, 204)
#         }
        
#     def load_json_data(self):
#         """Load JSON data"""
#         with open(self.json_file_path, 'r', encoding='utf-8') as file:
#             return json.load(file)
    
#     def create_presentation(self, output_path):
#         """Main method - creates presentation with dynamic layouts based on image availability"""
#         self.prs = Presentation()
#         self.slide_width = Inches(13.33)
#         self.slide_height = Inches(7.5)
#         self.prs.slide_width = self.slide_width
#         self.prs.slide_height = self.slide_height
        
#         print("🎨 Creating Light Theme Dynamic Layout Presentation")
#         print("🔄 Layout Strategy:")
#         print("   • 0 images → Text-only centered layout")
#         print("   • 1 image  → Single image on LEFT, text on RIGHT") 
#         print("   • 2+ images → Double images on RIGHT, text on LEFT")
        
#         # Create title slide
#         self.create_modern_title_slide()
        
#         # Create all content slides with DYNAMIC layouts
#         for slide_data in self.presentation_data["slides"]:
#             # Skip the first slide if it's the introduction
#             if slide_data["slide_number"] == 1:
#                 continue
                
#             available_images = self.get_available_images(slide_data)
#             image_count = len(available_images)
            
#             print(f"   Slide {slide_data['slide_number']}: '{slide_data['slide_title']}' - {image_count} image(s) available")
            
#             if image_count == 0:
#                 self.create_text_only_slide(slide_data)
#                 print(f"     → Using: Text-Only Layout")
#             elif image_count == 1:
#                 self.create_single_image_slide(slide_data, available_images[0])
#                 print(f"     → Using: Single Image Layout (Image LEFT)")
#             elif image_count >= 2:
#                 self.create_double_images_slide(slide_data, available_images[:2])
#                 print(f"     → Using: Double Images Layout (Images RIGHT)")
        
#         self.prs.save(output_path)
#         total_content_slides = len([s for s in self.presentation_data["slides"] if s["slide_number"] != 1])
#         print(f"✅ Light Theme Dynamic Layout Presentation created with {total_content_slides + 1} total slides!")
    
#     def get_available_images(self, slide_data):
#         """Get list of available images that actually exist"""
#         available = []
#         for img in slide_data.get("suggested_images", []):
#             img_path = img.get("image_path", "")
#             if img_path and os.path.exists(img_path):
#                 available.append(img)
#         return available
    
#     def create_modern_title_slide(self):
#         """Modern title slide with light theme - FIXED LAYOUT"""
#         slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
#         # Set white background
#         background = slide.background
#         fill = background.fill
#         fill.solid()
#         fill.fore_color.rgb = self.colors['background']
        
#         # Add geometric background elements
#         self.add_geometric_background(slide)
        
#         # Main title - FIXED: Smaller font and better positioning
#         title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), 
#                                            self.slide_width - Inches(1.0), Inches(2.0))  # Increased height
#         title_frame = title_box.text_frame
#         title_frame.text = self.presentation_data["presentation_title"]
#         title_frame.word_wrap = True
        
#         for paragraph in title_frame.paragraphs:
#             paragraph.font.size = Pt(36)  # Reduced from 44 to 36
#             paragraph.font.color.rgb = self.colors['text_primary']
#             paragraph.font.bold = True
#             paragraph.font.name = "Segoe UI"
#             paragraph.alignment = PP_ALIGN.CENTER
#             paragraph.space_after = Pt(0)
        
#         # Subtitle with accent - FIXED: Better positioning
#         subtitle_box = slide.shapes.add_textbox(Inches(2.0), Inches(4.5), 
#                                               self.slide_width - Inches(4.0), Inches(0.8))
#         subtitle_frame = subtitle_box.text_frame
#         subtitle_frame.text = "Advanced AI Solutions for Manufacturing Excellence"
        
#         for paragraph in subtitle_frame.paragraphs:
#             paragraph.font.size = Pt(18)  # Reduced from 20 to 18
#             paragraph.font.color.rgb = self.colors['primary']
#             paragraph.font.name = "Segoe UI"
#             paragraph.alignment = PP_ALIGN.CENTER
        
#         # Decorative elements
#         self.add_title_decoration(slide)
    
#     def create_text_only_slide(self, slide_data):
#         """Layout with only text content (centered) - FIXED: Centered text with proper gaps"""
#         slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
#         # Set white background
#         background = slide.background
#         fill = background.fill
#         fill.solid()
#         fill.fore_color.rgb = self.colors['background']
        
#         # Title
#         title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
#                                            self.slide_width - Inches(1.0), Inches(0.8))
#         title_frame = title_box.text_frame
#         title_frame.text = slide_data["slide_title"]
#         self.style_slide_title(title_frame)
        
#         # Centered text container - FIXED: Better dimensions
#         text_width = Inches(10.0)
#         text_left = (self.slide_width - text_width) / 2
#         text_container = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             text_left, Inches(1.5),
#             text_width, Inches(5.0)
#         )
#         text_container.fill.solid()
#         text_container.fill.fore_color.rgb = self.colors['card_bg']
#         text_container.line.color.rgb = self.colors['primary']
#         text_container.line.width = Pt(2)
        
#         # Text content with centered alignment - FIXED: Centered text
#         text_box = slide.shapes.add_textbox(
#             text_left + Inches(0.8), Inches(1.8),
#             text_width - Inches(1.6), Inches(4.4)
#         )
#         text_frame = text_box.text_frame
#         text_frame.word_wrap = True
#         text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # ✅ Now MSO_ANCHOR is defined
        
#         # Add all content points with centered alignment
#         content_items = slide_data.get("content", [])
#         for i, content_item in enumerate(content_items):
#             if i == 0:
#                 p = text_frame.paragraphs[0]
#             else:
#                 p = text_frame.add_paragraph()
            
#             # FIXED: Only one bullet, not double bullets
#             clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
#             p.text = "• " + clean_content  # Single bullet only
#             p.font.size = Pt(16)
#             p.font.color.rgb = self.colors['text_primary']
#             p.font.name = "Segoe UI"
#             p.space_after = Pt(12)  # Good gap between bullets
#             p.space_before = Pt(6)  # Space before each paragraph
#             p.alignment = PP_ALIGN.CENTER  # Center aligned text
#             p.line_spacing = 1.4
        
#         self.add_slide_number(slide, slide_data["slide_number"])
    
#     def create_single_image_slide(self, slide_data, image_data):
#         """Single image on LEFT side, text on RIGHT - FIXED: Text box boundaries"""
#         slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
#         # Set white background
#         background = slide.background
#         fill = background.fill
#         fill.solid()
#         fill.fore_color.rgb = self.colors['background']
        
#         # Title
#         title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
#                                            self.slide_width - Inches(1.0), Inches(0.8))
#         title_frame = title_box.text_frame
#         title_frame.text = slide_data["slide_title"]
#         self.style_slide_title(title_frame)
        
#         # Image Panel (LEFT) - 40% width
#         image_panel_width = self.slide_width * 0.40
#         image_panel = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             Inches(0.5), Inches(1.3),
#             image_panel_width, Inches(5.2)
#         )
#         image_panel.fill.solid()
#         image_panel.fill.fore_color.rgb = RGBColor(245, 245, 245)
#         image_panel.line.fill.background()
        
#         # Text Panel (RIGHT) - 55% width - FIXED: Proper boundaries
#         text_panel_width = self.slide_width * 0.55
#         text_panel_left = Inches(0.5) + image_panel_width + Inches(0.2)  # Reduced gap
#         text_panel = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             text_panel_left, Inches(1.3),
#             text_panel_width - Inches(0.2), Inches(5.2)  # Reduced width to fit
#         )
#         text_panel.fill.solid()
#         text_panel.fill.fore_color.rgb = self.colors['card_bg']
#         text_panel.line.color.rgb = self.colors['primary']
#         text_panel.line.width = Pt(2)
        
#         # Add image to LEFT panel
#         image_width = image_panel_width - Inches(1.2)
#         image_height = Inches(3.0)
#         image_left = Inches(0.5) + (image_panel_width - image_width) / 2
#         image_top = Inches(2.0)
        
#         try:
#             img = slide.shapes.add_picture(
#                 image_data["image_path"],
#                 image_left, image_top,
#                 width=image_width,
#                 height=image_height
#             )
#             # Add border to image
#             img.line.color.rgb = self.colors['image_border']
#             img.line.width = Pt(2)
#         except Exception as e:
#             print(f"🖼️ Image error: {e}")
#             self.add_image_placeholder(slide, image_left, image_top, image_width, image_height)
        
#         # Add text content to RIGHT panel - FIXED: Proper boundaries
#         text_box = slide.shapes.add_textbox(
#             text_panel_left + Inches(0.3), Inches(1.8),  # Reduced padding
#             text_panel_width - Inches(0.8), Inches(4.2)   # Proper width
#         )
#         text_frame = text_box.text_frame
#         text_frame.word_wrap = True
        
#         content_items = slide_data.get("content", [])
#         for i, content_item in enumerate(content_items):
#             if i == 0:
#                 p = text_frame.paragraphs[0]
#             else:
#                 p = text_frame.add_paragraph()
            
#             # FIXED: Only one bullet, not double bullets
#             clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
#             p.text = "• " + clean_content  # Single bullet only
#             p.font.size = Pt(14)
#             p.font.color.rgb = self.colors['text_primary']
#             p.font.name = "Segoe UI"
#             p.space_after = Pt(8)
#             p.line_spacing = 1.3
        
#         self.add_slide_number(slide, slide_data["slide_number"])
    
#     def create_double_images_slide(self, slide_data, image_data_list):
#         """Double images on RIGHT side, text on LEFT - FIXED: Text box boundaries"""
#         slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
#         # Set white background
#         background = slide.background
#         fill = background.fill
#         fill.solid()
#         fill.fore_color.rgb = self.colors['background']
        
#         # Title
#         title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
#                                            self.slide_width - Inches(1.0), Inches(0.8))
#         title_frame = title_box.text_frame
#         title_frame.text = slide_data["slide_title"]
#         self.style_slide_title(title_frame)
        
#         # Text Panel (LEFT) - 55% width - FIXED: Proper boundaries
#         text_panel_width = self.slide_width * 0.55
#         text_panel = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             Inches(0.5), Inches(1.3),
#             text_panel_width, Inches(5.2)
#         )
#         text_panel.fill.solid()
#         text_panel.fill.fore_color.rgb = self.colors['card_bg']
#         text_panel.line.color.rgb = self.colors['primary']
#         text_panel.line.width = Pt(2)
        
#         # Images Panel (RIGHT) - 40% width
#         images_panel_width = self.slide_width * 0.40
#         images_panel_left = Inches(0.5) + text_panel_width + Inches(0.2)  # Reduced gap
#         images_panel = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             images_panel_left, Inches(1.3),
#             images_panel_width - Inches(0.2), Inches(5.2)  # Reduced width to fit
#         )
#         images_panel.fill.solid()
#         images_panel.fill.fore_color.rgb = RGBColor(245, 245, 245)
#         images_panel.line.fill.background()
        
#         # Add text content to LEFT panel - FIXED: Proper boundaries
#         text_box = slide.shapes.add_textbox(
#             Inches(0.8), Inches(1.8),
#             text_panel_width - Inches(1.0), Inches(4.2)  # Reduced width to fit
#         )
#         text_frame = text_box.text_frame
#         text_frame.word_wrap = True
        
#         content_items = slide_data.get("content", [])
#         for i, content_item in enumerate(content_items):
#             if i == 0:
#                 p = text_frame.paragraphs[0]
#             else:
#                 p = text_frame.add_paragraph()
            
#             # FIXED: Only one bullet, not double bullets
#             clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
#             p.text = "• " + clean_content  # Single bullet only
#             p.font.size = Pt(14)
#             p.font.color.rgb = self.colors['text_primary']
#             p.font.name = "Segoe UI"
#             p.space_after = Pt(8)
#             p.line_spacing = 1.3
        
#         # Add two images to RIGHT panel (stacked vertically)
#         image_width = Inches(3.5)
#         image_height = Inches(1.8)
#         image_left = images_panel_left + (images_panel_width - image_width) / 2
        
#         # Add first image (TOP)
#         image1_top = Inches(1.8)
#         try:
#             img1 = slide.shapes.add_picture(
#                 image_data_list[0]["image_path"],
#                 image_left, image1_top,
#                 width=image_width,
#                 height=image_height
#             )
#             img1.line.color.rgb = self.colors['image_border']
#             img1.line.width = Pt(2)
#         except Exception as e:
#             print(f"🖼️ Image 1 error: {e}")
#             self.add_image_placeholder(slide, image_left, image1_top, image_width, image_height)
        
#         # Add second image (BOTTOM)
#         image2_top = image1_top + image_height + Inches(0.4)
#         try:
#             img2 = slide.shapes.add_picture(
#                 image_data_list[1]["image_path"],
#                 image_left, image2_top,
#                 width=image_width,
#                 height=image_height
#             )
#             img2.line.color.rgb = self.colors['image_border']
#             img2.line.width = Pt(2)
#         except Exception as e:
#             print(f"🖼️ Image 2 error: {e}")
#             self.add_image_placeholder(slide, image_left, image2_top, image_width, image_height)
        
#         self.add_slide_number(slide, slide_data["slide_number"])
    
#     def add_image_placeholder(self, slide, left, top, width, height):
#         """Add attractive image placeholder"""
#         placeholder = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE, left, top, width, height
#         )
#         placeholder.fill.solid()
#         placeholder.fill.fore_color.rgb = RGBColor(240, 240, 240)
#         placeholder.line.fill.background()
        
#         # Icon
#         icon_box = slide.shapes.add_textbox(
#             left + width/2 - Inches(0.3), top + height/2 - Inches(0.4),
#             Inches(0.6), Inches(0.6)
#         )
#         icon_frame = icon_box.text_frame
#         icon_frame.text = "🖼️"
#         icon_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
#         icon_frame.paragraphs[0].font.size = Pt(24)
#         icon_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
        
#         # Placeholder text
#         text_box = slide.shapes.add_textbox(
#             left, top + height/2 + Inches(0.2),
#             width, Inches(0.3)
#         )
#         text_frame = text_box.text_frame
#         text_frame.text = "Visual Content"
#         text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
#         text_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
#         text_frame.paragraphs[0].font.size = Pt(12)
#         text_frame.paragraphs[0].font.name = "Segoe UI"
    
#     # Background and styling methods
#     def add_geometric_background(self, slide):
#         """Add geometric elements to title slide background"""
#         for i in range(3):
#             circle = slide.shapes.add_shape(
#                 MSO_SHAPE.OVAL,
#                 Inches(2 + i*3), Inches(1 + i*0.5),
#                 Inches(1.2), Inches(1.2)
#             )
#             circle.fill.solid()
#             circle.fill.fore_color.rgb = RGBColor(245, 245, 245)
#             circle.line.fill.background()
    
#     def add_title_decoration(self, slide):
#         """Add decorative elements to title slide"""
#         line = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             Inches(3.0), self.slide_height - Inches(0.8),
#             self.slide_width - Inches(6.0), Inches(0.03)
#         )
#         line.fill.solid()
#         line.fill.fore_color.rgb = self.colors['primary']
#         line.line.fill.background()
    
#     def style_slide_title(self, title_frame):
#         """Style slide titles"""
#         for paragraph in title_frame.paragraphs:
#             paragraph.font.size = Pt(24)
#             paragraph.font.bold = True
#             paragraph.font.color.rgb = self.colors['text_primary']
#             paragraph.font.name = "Segoe UI"
#             paragraph.alignment = PP_ALIGN.CENTER
    
#     def add_slide_number(self, slide, slide_number):
#         """Add slide number"""
#         total_slides = len([s for s in self.presentation_data["slides"] if s["slide_number"] != 1]) + 1
        
#         number_bg = slide.shapes.add_shape(
#             MSO_SHAPE.ROUNDED_RECTANGLE,
#             self.slide_width - Inches(1.8), self.slide_height - Inches(0.5),
#             Inches(1.3), Inches(0.35)
#         )
#         number_bg.fill.solid()
#         number_bg.fill.fore_color.rgb = RGBColor(248, 249, 250)
#         number_bg.line.color.rgb = self.colors['border']
#         number_bg.line.width = Pt(1)
        
#         number_box = slide.shapes.add_textbox(
#             self.slide_width - Inches(1.8), self.slide_height - Inches(0.5),
#             Inches(1.3), Inches(0.35)
#         )
#         text_frame = number_box.text_frame
#         text_frame.text = f"{slide_number}/{total_slides}"
#         text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
#         text_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
#         text_frame.paragraphs[0].font.size = Pt(12)
#         text_frame.paragraphs[0].font.name = "Segoe UI"





from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import json
import os

class JSONAdapter:
    @staticmethod
    def adapt_json_for_ppt(input_json_path, output_json_path):
        """Adapt the image retrieval JSON structure to match PPT generator expectations"""
        with open(input_json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Transform the structure
        adapted_data = {
            "presentation_title": data.get("presentation_title", "AI Presentation"),
            "slides": []
        }
        
        for slide in data.get("slides", []):
            # Convert content to list format expected by PPT generator
            content_list = []
            if slide.get("content"):
                content_text = slide["content"]
                if isinstance(content_text, str):
                    # Split content by sentences or newlines for bullet points
                    sentences = [s.strip() for s in content_text.split('\n') if s.strip()]
                    content_list = sentences[:6]  # Limit to 6 bullet points
                else:
                    content_list = content_text
            
            # Adapt images structure
            suggested_images = []
            for img in slide.get("images", []):
                suggested_images.append({
                    "image_path": img.get("image_path", ""),
                    "caption": img.get("caption", "Image"),
                    "similarity_score": img.get("similarity_score", 0)
                })
            
            adapted_slide = {
                "slide_number": slide.get("slide_number", len(adapted_data["slides"]) + 1),
                "slide_title": slide.get("slide_title", "Untitled Slide"),
                "content": content_list,
                "suggested_images": suggested_images
            }
            
            adapted_data["slides"].append(adapted_slide)
        
        # Save adapted JSON
        with open(output_json_path, 'w', encoding='utf-8') as f:
            json.dump(adapted_data, f, indent=2)
        
        print(f"✅ JSON adapted successfully for template2!")
        return adapted_data

class DynamicLayoutPresentationGenerator:
    def __init__(self, json_file_path):
        self.json_file_path = json_file_path
        self.presentation_data = self.load_json_data()
        self.prs = None
        
        # 🌤️ LIGHT THEME COLORS
        self.colors = {
            'background': RGBColor(255, 255, 255),
            'card_bg': RGBColor(248, 249, 250),
            'primary': RGBColor(0, 102, 204),
            'accent': RGBColor(255, 87, 34),
            'success': RGBColor(46, 204, 113),
            'text_primary': RGBColor(33, 37, 41),
            'text_secondary': RGBColor(108, 117, 125),
            'border': RGBColor(222, 226, 230),
            'image_border': RGBColor(0, 102, 204)
        }
        
    def load_json_data(self):
        """Load JSON data"""
        with open(self.json_file_path, 'r', encoding='utf-8') as file:
            return json.load(file)
    
    def create_presentation(self, output_path):
        """Main method - creates presentation with dynamic layouts based on image availability"""
        self.prs = Presentation()
        self.slide_width = Inches(13.33)
        self.slide_height = Inches(7.5)
        self.prs.slide_width = self.slide_width
        self.prs.slide_height = self.slide_height
        
        print("🎨 Creating Light Theme Dynamic Layout Presentation")
        print("🔄 Layout Strategy:")
        print("   • 0 images → Text-only centered layout")
        print("   • 1 image  → Single image on LEFT, text on RIGHT") 
        print("   • 2+ images → Double images on RIGHT, text on LEFT")
        
        # Create title slide
        self.create_modern_title_slide()
        
        # Create all content slides with DYNAMIC layouts
        for slide_data in self.presentation_data["slides"]:
            # Skip the first slide if it's the introduction
            if slide_data["slide_number"] == 1:
                continue
                
            available_images = self.get_available_images(slide_data)
            image_count = len(available_images)
            
            print(f"   Slide {slide_data['slide_number']}: '{slide_data['slide_title']}' - {image_count} image(s) available")
            
            if image_count == 0:
                self.create_text_only_slide(slide_data)
                print(f"     → Using: Text-Only Layout")
            elif image_count == 1:
                self.create_single_image_slide(slide_data, available_images[0])
                print(f"     → Using: Single Image Layout (Image LEFT)")
            elif image_count >= 2:
                self.create_double_images_slide(slide_data, available_images[:2])
                print(f"     → Using: Double Images Layout (Images RIGHT)")
        
        self.prs.save(output_path)
        total_content_slides = len([s for s in self.presentation_data["slides"] if s["slide_number"] != 1])
        print(f"✅ Light Theme Dynamic Layout Presentation created with {total_content_slides + 1} total slides!")
    
    def get_available_images(self, slide_data):
        """Get list of available images that actually exist"""
        available = []
        for img in slide_data.get("suggested_images", []):
            img_path = img.get("image_path", "")
            if img_path and os.path.exists(img_path):
                available.append(img)
        return available
    
    def create_modern_title_slide(self):
        """Modern title slide with light theme - FIXED LAYOUT"""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
        # Set white background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.colors['background']
        
        # Add geometric background elements
        self.add_geometric_background(slide)
        
        # Main title - FIXED: Smaller font and better positioning
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), 
                                           self.slide_width - Inches(1.0), Inches(2.0))  # Increased height
        title_frame = title_box.text_frame
        title_frame.text = self.presentation_data["presentation_title"]
        title_frame.word_wrap = True
        
        for paragraph in title_frame.paragraphs:
            paragraph.font.size = Pt(36)  # Reduced from 44 to 36
            paragraph.font.color.rgb = self.colors['text_primary']
            paragraph.font.bold = True
            paragraph.font.name = "Segoe UI"
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.space_after = Pt(0)
        
        # Subtitle with accent - FIXED: Better positioning
        subtitle_box = slide.shapes.add_textbox(Inches(2.0), Inches(4.5), 
                                              self.slide_width - Inches(4.0), Inches(0.8))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = "Advanced AI Solutions for Manufacturing Excellence"
        
        for paragraph in subtitle_frame.paragraphs:
            paragraph.font.size = Pt(18)  # Reduced from 20 to 18
            paragraph.font.color.rgb = self.colors['primary']
            paragraph.font.name = "Segoe UI"
            paragraph.alignment = PP_ALIGN.CENTER
        
        # Decorative elements
        self.add_title_decoration(slide)
    
    def create_text_only_slide(self, slide_data):
        """Layout with only text content (centered) - FIXED: Centered text with proper gaps"""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
        # Set white background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.colors['background']
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
                                           self.slide_width - Inches(1.0), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = slide_data["slide_title"]
        self.style_slide_title(title_frame)
        
        # Centered text container - FIXED: Better dimensions
        text_width = Inches(10.0)
        text_left = (self.slide_width - text_width) / 2
        text_container = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            text_left, Inches(1.5),
            text_width, Inches(5.0)
        )
        text_container.fill.solid()
        text_container.fill.fore_color.rgb = self.colors['card_bg']
        text_container.line.color.rgb = self.colors['primary']
        text_container.line.width = Pt(2)
        
        # Text content with centered alignment - FIXED: Centered text
        text_box = slide.shapes.add_textbox(
            text_left + Inches(0.8), Inches(1.8),
            text_width - Inches(1.6), Inches(4.4)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # Center vertically
        
        # Add all content points with centered alignment
        content_items = slide_data.get("content", [])
        for i, content_item in enumerate(content_items):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            
            # FIXED: Only one bullet, not double bullets
            clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
            p.text = "• " + clean_content  # Single bullet only
            p.font.size = Pt(18)  # Increased from 16
            p.font.color.rgb = self.colors['text_primary']
            p.font.name = "Segoe UI"
            p.space_after = Pt(12)  # Good gap between bullets
            p.space_before = Pt(6)  # Space before each paragraph
            p.alignment = PP_ALIGN.CENTER  # Center aligned text
            p.line_spacing = 1.4
        
        self.add_slide_number(slide, slide_data["slide_number"])
    
    def create_single_image_slide(self, slide_data, image_data):
        """Single image on LEFT side, text on RIGHT - FIXED: Centered text"""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
        # Set white background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.colors['background']
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
                                           self.slide_width - Inches(1.0), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = slide_data["slide_title"]
        self.style_slide_title(title_frame)
        
        # Image Panel (LEFT) - 40% width
        image_panel_width = self.slide_width * 0.40
        image_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5), Inches(1.3),
            image_panel_width, Inches(5.2)
        )
        image_panel.fill.solid()
        image_panel.fill.fore_color.rgb = RGBColor(245, 245, 245)
        image_panel.line.fill.background()
        
        # Text Panel (RIGHT) - 55% width - FIXED: Proper boundaries
        text_panel_width = self.slide_width * 0.55
        text_panel_left = Inches(0.5) + image_panel_width + Inches(0.2)  # Reduced gap
        text_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            text_panel_left, Inches(1.3),
            text_panel_width - Inches(0.2), Inches(5.2)  # Reduced width to fit
        )
        text_panel.fill.solid()
        text_panel.fill.fore_color.rgb = self.colors['card_bg']
        text_panel.line.color.rgb = self.colors['primary']
        text_panel.line.width = Pt(2)
        
        # Add image to LEFT panel
        image_width = image_panel_width - Inches(1.2)
        image_height = Inches(3.0)
        image_left = Inches(0.5) + (image_panel_width - image_width) / 2
        image_top = Inches(2.0)
        
        try:
            img = slide.shapes.add_picture(
                image_data["image_path"],
                image_left, image_top,
                width=image_width,
                height=image_height
            )
            # Add border to image
            img.line.color.rgb = self.colors['image_border']
            img.line.width = Pt(2)
        except Exception as e:
            print(f"🖼️ Image error: {e}")
            self.add_image_placeholder(slide, image_left, image_top, image_width, image_height)
        
        # Add text content to RIGHT panel - FIXED: Centered text
        text_box = slide.shapes.add_textbox(
            text_panel_left + Inches(0.3), Inches(1.8),  # Reduced padding
            text_panel_width - Inches(0.8), Inches(4.2)   # Proper width
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # ✅ Center text vertically
        
        content_items = slide_data.get("content", [])
        for i, content_item in enumerate(content_items):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            
            # FIXED: Only one bullet, not double bullets
            clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
            p.text = "• " + clean_content  # Single bullet only
            p.font.size = Pt(16)  # Increased from 14
            p.font.color.rgb = self.colors['text_primary']
            p.font.name = "Segoe UI"
            p.space_after = Pt(10)  # Increased spacing
            p.space_before = Pt(5)  # Added space before
            p.alignment = PP_ALIGN.CENTER  # ✅ Center align text
            p.line_spacing = 1.3
        
        self.add_slide_number(slide, slide_data["slide_number"])
    
    def create_double_images_slide(self, slide_data, image_data_list):
        """Double images on RIGHT side, text on LEFT - FIXED: Centered text"""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
        # Set white background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.colors['background']
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), 
                                           self.slide_width - Inches(1.0), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = slide_data["slide_title"]
        self.style_slide_title(title_frame)
        
        # Text Panel (LEFT) - 55% width - FIXED: Proper boundaries
        text_panel_width = self.slide_width * 0.55
        text_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5), Inches(1.3),
            text_panel_width, Inches(5.2)
        )
        text_panel.fill.solid()
        text_panel.fill.fore_color.rgb = self.colors['card_bg']
        text_panel.line.color.rgb = self.colors['primary']
        text_panel.line.width = Pt(2)
        
        # Images Panel (RIGHT) - 40% width
        images_panel_width = self.slide_width * 0.40
        images_panel_left = Inches(0.5) + text_panel_width + Inches(0.2)  # Reduced gap
        images_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            images_panel_left, Inches(1.3),
            images_panel_width - Inches(0.2), Inches(5.2)  # Reduced width to fit
        )
        images_panel.fill.solid()
        images_panel.fill.fore_color.rgb = RGBColor(245, 245, 245)
        images_panel.line.fill.background()
        
        # Add text content to LEFT panel - FIXED: Centered text
        text_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.8),
            text_panel_width - Inches(1.0), Inches(4.2)  # Reduced width to fit
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # ✅ Center text vertically
        
        content_items = slide_data.get("content", [])
        for i, content_item in enumerate(content_items):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            
            # FIXED: Only one bullet, not double bullets
            clean_content = str(content_item).replace('•', '').replace('▸', '').strip()
            p.text = "• " + clean_content  # Single bullet only
            p.font.size = Pt(16)  # Increased from 14
            p.font.color.rgb = self.colors['text_primary']
            p.font.name = "Segoe UI"
            p.space_after = Pt(10)  # Increased spacing
            p.space_before = Pt(5)  # Added space before
            p.alignment = PP_ALIGN.CENTER  # ✅ Center align text
            p.line_spacing = 1.3
        
        # Add two images to RIGHT panel (stacked vertically)
        image_width = Inches(3.5)
        image_height = Inches(1.8)
        image_left = images_panel_left + (images_panel_width - image_width) / 2
        
        # Add first image (TOP)
        image1_top = Inches(1.8)
        try:
            img1 = slide.shapes.add_picture(
                image_data_list[0]["image_path"],
                image_left, image1_top,
                width=image_width,
                height=image_height
            )
            img1.line.color.rgb = self.colors['image_border']
            img1.line.width = Pt(2)
        except Exception as e:
            print(f"🖼️ Image 1 error: {e}")
            self.add_image_placeholder(slide, image_left, image1_top, image_width, image_height)
        
        # Add second image (BOTTOM)
        image2_top = image1_top + image_height + Inches(0.4)
        try:
            img2 = slide.shapes.add_picture(
                image_data_list[1]["image_path"],
                image_left, image2_top,
                width=image_width,
                height=image_height
            )
            img2.line.color.rgb = self.colors['image_border']
            img2.line.width = Pt(2)
        except Exception as e:
            print(f"🖼️ Image 2 error: {e}")
            self.add_image_placeholder(slide, image_left, image2_top, image_width, image_height)
        
        self.add_slide_number(slide, slide_data["slide_number"])
    
    def add_image_placeholder(self, slide, left, top, width, height):
        """Add attractive image placeholder"""
        placeholder = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, width, height
        )
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(240, 240, 240)
        placeholder.line.fill.background()
        
        # Icon
        icon_box = slide.shapes.add_textbox(
            left + width/2 - Inches(0.3), top + height/2 - Inches(0.4),
            Inches(0.6), Inches(0.6)
        )
        icon_frame = icon_box.text_frame
        icon_frame.text = "🖼️"
        icon_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        icon_frame.paragraphs[0].font.size = Pt(24)
        icon_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
        
        # Placeholder text
        text_box = slide.shapes.add_textbox(
            left, top + height/2 + Inches(0.2),
            width, Inches(0.3)
        )
        text_frame = text_box.text_frame
        text_frame.text = "Visual Content"
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
        text_frame.paragraphs[0].font.size = Pt(12)
        text_frame.paragraphs[0].font.name = "Segoe UI"
    
    # Background and styling methods
    def add_geometric_background(self, slide):
        """Add geometric elements to title slide background"""
        for i in range(3):
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(2 + i*3), Inches(1 + i*0.5),
                Inches(1.2), Inches(1.2)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = RGBColor(245, 245, 245)
            circle.line.fill.background()
    
    def add_title_decoration(self, slide):
        """Add decorative elements to title slide"""
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(3.0), self.slide_height - Inches(0.8),
            self.slide_width - Inches(6.0), Inches(0.03)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.colors['primary']
        line.line.fill.background()
    
    def style_slide_title(self, title_frame):
        """Style slide titles"""
        for paragraph in title_frame.paragraphs:
            paragraph.font.size = Pt(24)
            paragraph.font.bold = True
            paragraph.font.color.rgb = self.colors['text_primary']
            paragraph.font.name = "Segoe UI"
            paragraph.alignment = PP_ALIGN.CENTER
    
    def add_slide_number(self, slide, slide_number):
        """Add slide number"""
        total_slides = len([s for s in self.presentation_data["slides"] if s["slide_number"] != 1]) + 1
        
        number_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            self.slide_width - Inches(1.8), self.slide_height - Inches(0.5),
            Inches(1.3), Inches(0.35)
        )
        number_bg.fill.solid()
        number_bg.fill.fore_color.rgb = RGBColor(248, 249, 250)
        number_bg.line.color.rgb = self.colors['border']
        number_bg.line.width = Pt(1)
        
        number_box = slide.shapes.add_textbox(
            self.slide_width - Inches(1.8), self.slide_height - Inches(0.5),
            Inches(1.3), Inches(0.35)
        )
        text_frame = number_box.text_frame
        text_frame.text = f"{slide_number}/{total_slides}"
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text_frame.paragraphs[0].font.color.rgb = self.colors['text_secondary']
        text_frame.paragraphs[0].font.size = Pt(12)
        text_frame.paragraphs[0].font.name = "Segoe UI"